import pandas as pd
import random
import json
from openpyxl import load_workbook
from datetime import datetime
from openpyxl.styles import Border, Side, Alignment, Font

with open('config.json', 'r',encoding='utf-8') as f:
    config = json.load(f)

file_path = config['file_path']
min_samples = config['min_samples']
max_samples = config['max_samples']
delimiter_samples = config['delimiter_samples']

dfs = {i: pd.read_excel(file_path, sheet_name=i) for i in range(13)}

def random_sample_with_delimiter(df, delimiter_column, num_samples):
    nan_indices = df[df[delimiter_column].isna()].index.tolist()

    groups = []
    start_idx = 0
    for idx in nan_indices:
        groups.append(df.iloc[start_idx:idx+1])  
        start_idx = idx + 1
    groups.append(df.iloc[start_idx:])  

    groups = [group for group in groups if not group.empty]

    sampled_groups = random.sample(groups, min(num_samples, len(groups)))

    sampled_df = pd.concat(sampled_groups, ignore_index=True)

    return sampled_df, len(sampled_groups)

def random_sample_within_range(group, min_samples, max_samples):
    num_samples = random.randint(min_samples, max_samples)
    return group.sample(n=min(num_samples, len(group)), random_state=1)

delimiter_dfs = {}
for i in range(config['sheet_processing']['delimiter_range'][0], config['sheet_processing']['delimiter_range'][1] + 1):
    delimiter_dfs[f'mul_conversation_{i-3}_df'], delimiter_dfs[f'mul_conversation_{i-3}_groups'] = random_sample_with_delimiter(dfs[i], '语料', delimiter_samples)

groupby_dfs = {}
for i in range(config['sheet_processing']['groupby_range'][0], config['sheet_processing']['groupby_range'][1] + 1):
    groupby_dfs[f'df_{i}'] = dfs[i].groupby('分类').apply(lambda x: random_sample_within_range(x, min_samples, max_samples)).reset_index(drop=True)

for i in range(7, 13):
    groupby_dfs[f'df_{i}'] = dfs[i].groupby('分类').apply(lambda x: random_sample_within_range(x, min_samples, max_samples)).reset_index(drop=True)

num_groups = {}
for df in groupby_dfs.values():
    num_groups.update(df['分类'].value_counts().to_dict())

num_groups.update({
    '垂域内多轮对话': delimiter_dfs['mul_conversation_1_groups'],
    '垂域间多轮对话': delimiter_dfs['mul_conversation_2_groups'],
    '跨外部模型多轮对话': delimiter_dfs['mul_conversation_3_groups']
})

current_date = datetime.now()
formatted_date = current_date.strftime('%m%d')
excel_name = f'LLM中控能力分类及标准{formatted_date}.xlsx'

with pd.ExcelWriter(excel_name) as writer:
    dfs[0].to_excel(writer, sheet_name='总览', index=False)
    
    sheet_names = ['意图理解', '意图引导', '实体抽取', '垂域内多轮对话', '垂域间多轮对话', '跨外部模型多轮对话', 
                   '安全性', '语义拒识', '内容搜索', '复杂任务', '自然对话生成', '话术多样性']
    
    for i, name in enumerate(sheet_names):
        if i < 3:
            groupby_dfs[f'df_{i+1}'].to_excel(writer, sheet_name=name, index=False)
        elif i < 6:
            delimiter_dfs[f'mul_conversation_{i-2}_df'].to_excel(writer, sheet_name=name, index=False)
        else:
            groupby_dfs[f'df_{i+1}'].to_excel(writer, sheet_name=name, index=False)

wb = load_workbook(excel_name)
ws = wb['总览']

ws.insert_cols(4)
ws.cell(row=1, column=4).value = 'Num'

for row in range(2, ws.max_row + 1):
    category = ws.cell(row=row, column=3).value
    if category in num_groups:
        ws.cell(row=row, column=4).value = num_groups[category]

def ws_width_save(sheet_name, column, width):
    ws = wb[sheet_name]
    ws.column_dimensions[column].width = width
    wb.save(excel_name)

header_cell = ws.cell(row=1, column=4)
header_cell.font = Font(bold=True)
header_cell.alignment = Alignment(horizontal='center', vertical='center')

def set_border(ws, cell_range):
    border = Border(left=Side(border_style='thin'),
                    right=Side(border_style='thin'),
                    top=Side(border_style='thin'),
                    bottom=Side(border_style='thin'))

    for row in ws[cell_range]:
        for cell in row:
            cell.border = border

ws_width_save('总览', 'B', 40)
ws_width_save('总览', 'C', 40)
ws_width_save('总览', 'D', 5)  
ws_width_save('总览', 'E', 110)
set_border(ws, f'A1:E{ws.max_row}')

wb.save(excel_name)

for sheet_name in sheet_names:
    if sheet_name in ['垂域内多轮对话', '垂域间多轮对话', '跨外部模型多轮对话', '自然对话生成']:
        ws_width_save(sheet_name, 'A', 20)
    ws_width_save(sheet_name, 'B', 40)

ws.column_dimensions['D'].width = 10
ws.column_dimensions['E'].width = 110

wb.save(excel_name)