import tkinter as tk
from tkinter import messagebox, filedialog
from tkinter import ttk
from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.options import Options
import configparser
import openpyxl
import re
import threading
import logging

# 配置日志记录
logging.basicConfig(filename='error.log', level=logging.ERROR)

# 获取音乐ID列表
def get_music_ids(driver):
    table = driver.find_element(By.TAG_NAME, "table")
    a_tags = table.find_elements(By.TAG_NAME, "a")
    hrefs = {a.get_attribute("href") for a in a_tags if "song?id=" in a.get_attribute("href")}
    musics_id = []
    for href in hrefs:
        match = re.search(r'song\?id=(\d+)', href)
        if match:
            musics_id.append(match.group(1))
    return musics_id

# 获取音乐名称
def get_music_name(driver):
    song_title_element = driver.find_element(By.XPATH, '//*[@class="tit"]/em')
    song_title = song_title_element.text
    music_name = re.sub(r'\([^)]*\)', '', song_title)  # 去除括号及其内容
    music_name = re.sub(r'[^\w\s]', '', music_name)  # 去除所有非字母、数字、空格的字符
    music_name = re.sub(r'\s+', ' ', music_name)  # 去除多余的空格
    return music_name.strip()  # 去除首尾空格

# 获取音乐作者
def get_music_author(driver):
    artist_element = driver.find_element(By.XPATH, '//*[@class="des s-fc4"]/span')
    title_value = artist_element.get_attribute('title').strip('title=')
    artist_name = re.sub(r'/', '和', title_value)
    # 紧凑排列人名
    # artist_name = re.sub(r'\s+', ' ', artist_name)  # 使用正则表达式替换多个空格为一个空格
    artist_name = re.sub(r'\s', '', artist_name)  # 使用正则表达式替换所有空格为空
    return artist_name

# 线程终止标志
stop_thread = False

# 显示歌单音乐列表
def show_music_list():
    global stop_thread
    stop_thread = False
    output_text.delete('1.0', tk.END)
    progress.start()
    try:
        driver.get(f"https://music.163.com/#/discover/toplist?id={playlist_id.get()}")
        driver.switch_to.frame("contentFrame")
        music_ids = get_music_ids(driver)
        music_listbox.delete(0, tk.END)  # 清空列表框
        for music_id in music_ids:
            if stop_thread:
                break
            driver.get(f'https://music.163.com/#/song?id={music_id}')
            driver.switch_to.frame("contentFrame")
            music_name = get_music_name(driver)
            author_name = get_music_author(driver)
            music_info = f"音乐：{music_name}-->作者：{author_name}-->音乐ID：{music_id}"
            music_listbox.insert(tk.END, music_info)
            output_text.insert(tk.END, f"{music_info}\n")
    except Exception as e:
        logging.error("Error in show_music_list: %s", e)
    finally:
        progress.stop()


# 异步显示歌单音乐列表
def show_music_list_async():
    thread = threading.Thread(target=show_music_list)
    thread.start()

# 终止获取歌单音乐列表
def stop_music_list():
    global stop_thread
    stop_thread = True
    output_text.insert(tk.END, "中止操作成功\n")

# 保存音乐信息到Excel
def save_to_excel(): 
    file_path = filedialog.askopenfilename(title="选择Excel文件", filetypes=[("Excel files", "*.xlsx")])
    if not file_path:
        return
    global stop_thread
    stop_thread = False
    output_text.delete('1.0', tk.END)
    progress.start()
    try:
        workbook = openpyxl.load_workbook(file_path)
        sheet = workbook.active
    except FileNotFoundError:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
    
    try:
        driver.get(f"https://music.163.com/discover/toplist?id={playlist_id.get()}")
        driver.switch_to.frame("contentFrame")
        music_ids = get_music_ids(driver)
        row = 1
        while sheet.cell(row=row, column=1).value:
            row += 1
        
        for music_id in music_ids:
            driver.get(f'https://music.163.com/song?id={music_id}')
            if stop_thread:
                break
            driver.switch_to.frame("contentFrame")
            music_name = get_music_name(driver)
            author_name = get_music_author(driver)
            
            sheet.cell(row=row, column=1).value = f"{music_name} --> {author_name}"
            output_text.insert(tk.END, f"{music_name} --> {author_name} 保存成功\n")
            row += 1
        
        workbook.save(file_path)
        messagebox.showinfo("保存成功", "音乐信息已成功保存到Excel文件")
    except Exception as e:
        logging.error("Error in save_to_excel: %s", e)
    finally:
        progress.stop()

#异步显示保存的音乐
def show_save_music_list_async():
    thread = threading.Thread(target=save_to_excel)
    thread.start()

def del_English_autohor():
    file_path = filedialog.askopenfilename(title="选择Excel文件", filetypes=[("Excel files", "*.xlsx")])
    global stop_thread
    stop_thread = False
    output_text.delete('1.0', tk.END)
    progress.start()
    try:
        workbook = openpyxl.load_workbook(file_path)
        sheet = workbook.active
    except FileNotFoundError:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        logging.error("文件未找到")
        return

    try:
        rows_to_delete = []
        for row_num, row in enumerate(sheet.iter_rows(min_row=2), 2):
            # 确保值是字符串类型
            a_value = str(row[0].value) if row[0].value is not None else None
            b_value = str(row[1].value) if row[1].value is not None else None

            # 检查A列是否为全英文或全非中文字符，如果是则删除整行
            if a_value and (re.match(r'^[a-zA-Z\s]+$', a_value) or not re.search(r'[\u4e00-\u9fa5]', a_value)):
                rows_to_delete.append(row_num)

        # 从后往前删除行，避免索引问题
        for row_num in reversed(rows_to_delete):
            sheet.delete_rows(row_num)
            output_text.insert(tk.END, f"删除第 {row_num} 行，A列值: {row[0].value}, B列值: {row[1].value}\n")

        workbook.save(file_path)
        output_text.insert(tk.END, "所有A列全英文或全非中文字符的作者已成功从Excel文件删除\n")
        messagebox.showinfo("删除成功", "A列全英文或全非中文字符的作者已成功从Excel文件删除")
    except Exception as e:
        logging.error("Error in del_to_excel: %s", e)
        output_text.insert(tk.END, f"处理过程中发生错误: {e}\n")
    finally:
        progress.stop()

def show_del_English_author():
    thread = threading.Thread(target=del_English_autohor)
    thread.start()

def return_to_main_window():
    function_window.destroy()  # 关闭功能窗口
    root.deiconify()  # 显示主窗口

# 初始化配置和Selenium
config = configparser.ConfigParser()
config.read('config.ini')
binary_location = config['chrome']['binary_location']
driver_path = config['chrome']['driver_path']
chrome_options = Options()
chrome_options.add_experimental_option("prefs", {
    "profile.managed_default_content_settings.images": 2  # 2表示阻止加载图片
})
chrome_options.add_argument('--enable-chrome-browser-cloud-management')
chrome_options.binary_location = binary_location
service = Service(driver_path)
driver = webdriver.Chrome(service=service, options=chrome_options)

# 创建主窗口
root = tk.Tk()
root.title("音乐歌单爬取")
root.geometry("800x600")

tk.Label(root, text="请输入歌单ID:").pack(pady=10)
playlist_id = tk.StringVar()
tk.Entry(root, textvariable=playlist_id).pack(pady=5)

def open_function_window():
    global function_window 
    root.withdraw()  # 隐藏主窗口

    function_window = tk.Toplevel(root)
    function_window.title("功能窗口")
    function_window.geometry("800x600")

    button_frame = tk.Frame(function_window)
    button_frame.pack(pady=10)

    tk.Button(button_frame, text="列出音乐", command=show_music_list_async).pack(side=tk.LEFT, padx=5)
    tk.Button(button_frame, text="终止", command=stop_music_list).pack(side=tk.LEFT, padx=5)
    tk.Button(button_frame, text="保存到Excel", command=show_save_music_list_async).pack(side=tk.LEFT, padx=5)
    tk.Button(button_frame, text="筛选英文作者删除",command=show_del_English_author).pack(side=tk.LEFT, padx=5)
    tk.Button(button_frame, text="退出", command=lambda: (driver.quit(), root.destroy())).pack(side=tk.LEFT, padx=5)
    tk.Button(button_frame, text="返回", command=return_to_main_window).pack(side=tk.LEFT, padx=5)

    global output_text
    output_text = tk.Text(function_window, height=20, width=70)
    output_text.pack(pady=10)

    global progress
    progress = ttk.Progressbar(function_window, mode='indeterminate')
    progress.pack(pady=10)

    global music_listbox
    music_listbox = tk.Listbox(function_window, height=10, width=70)
    music_listbox.pack(pady=10)

    # 添加列表框选择事件处理函数
    music_listbox.bind('<<ListboxSelect>>', on_music_select)

    function_window.protocol("WM_DELETE_WINDOW", lambda: (driver.quit(), root.destroy()))

tk.Button(root, text="确定", command=open_function_window).pack(pady=10)

# 列表框选择事件处理函数
def on_music_select(event):
    selected_index = music_listbox.curselection()
    if selected_index:
        selected_music = music_listbox.get(selected_index)
        messagebox.showinfo("选择的歌曲", selected_music)

root.mainloop()

driver.quit()