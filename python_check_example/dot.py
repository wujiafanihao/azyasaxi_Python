import openpyxl
import re

input_file = "douban_answer.xlsx"
sheet_index = 0

input_wb = openpyxl.load_workbook(input_file)
input_ws = input_wb.worksheets[sheet_index]

punctuation_pattern = re.compile(r'[^\w\s]')

current_row = 1
while current_row <= input_ws.max_row:
    check_1_cell = input_ws.cell(row=current_row, column=1).value
    check_2_cell = input_ws.cell(row=current_row, column=2).value
    check_3_cell = input_ws.cell(row=current_row, column=3).value
    
    if (check_2_cell and punctuation_pattern.search(check_2_cell)) or (check_3_cell and punctuation_pattern.search(check_3_cell)):
    # if(check_1_cell and punctuation_pattern.search(check_1_cell)):
        input_ws.delete_rows(current_row, 1)
        print(f"删除第 {current_row} 行")
    else:
        current_row += 1

input_wb.save(input_file)
