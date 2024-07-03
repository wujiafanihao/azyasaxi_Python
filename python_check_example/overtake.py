import openpyxl

input_file = "唯一片名.xlsx"
sheet_index = 1

input_wb = openpyxl.load_workbook(input_file)
input_ws = input_wb.worksheets[sheet_index]

current_row = 1

while current_row <= input_ws.max_row:
    check_1_cell = input_ws.cell(row=current_row, column=1).value
    check_2_cell = input_ws.cell(row=current_row, column=2).value
    check_3_cell = input_ws.cell(row=current_row, column=3).value

    if check_1_cell and check_2_cell and check_3_cell:
        if check_3_cell == "多字" or check_3_cell == "重复":
            if len(check_2_cell) <= len(check_1_cell):
                input_ws.delete_rows(current_row)
                print(f"删除第 {current_row} 行")
                continue  
    current_row += 1

input_wb.save(input_file)
