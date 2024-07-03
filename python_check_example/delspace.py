from openpyxl import load_workbook

# 加载Excel文件
wb = load_workbook('douban_answer.xlsx')
ws = wb.active

# 从下到上检查每一行，如果是空行就删除
max_row = ws.max_row
for row in range(max_row, 0, -1):
    if all([cell.value is None for cell in ws[row]]):
        ws.delete_rows(row)

# 保存修改后的文件
wb.save('douban_answer.xlsx')
