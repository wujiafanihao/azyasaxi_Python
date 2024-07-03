import openpyxl

input_file = "唯一片名.xlsx"

input_wb = openpyxl.load_workbook(input_file)
input_ws = input_wb.active

current_row = 1

while current_row <= input_ws.max_row:
    check_1_cell = input_ws.cell(row=current_row, column=1).value
    check_2_cell = input_ws.cell(row=current_row, column=2).value
    check_3_cell = input_ws.cell(row=current_row, column=3).value

    if check_1_cell and check_2_cell and check_3_cell:
        if check_3_cell == "实体颠倒":
            if check_1_cell == check_2_cell[::-1]:
                input_ws.delete_rows(current_row)
                print(f"删除第 {current_row} 行")
                print(check_2_cell)
                print(check_2_cell[::-1])
                continue  
    current_row += 1

input_wb.save(input_file)
