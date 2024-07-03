from dotenv import load_dotenv
load_dotenv()
from langchain_community.chat_models import ChatOpenAI
from langchain.memory import ConversationSummaryBufferMemory
from langchain_core.prompts import ChatPromptTemplate,SystemMessagePromptTemplate,HumanMessagePromptTemplate,MessagesPlaceholder
from langchain.chains.conversation.base import ConversationChain

import openpyxl
import os
import re

input_file = "LLAMA3.xlsx"
output_file = "TEST.xlsx"
current_row = 2
sheet_index = 10
input_wb = openpyxl.load_workbook(input_file)
input_ws = input_wb.worksheets[sheet_index]
title_cell = input_ws.cell(row=2, column=1)
Model = input_ws.cell(row=current_row, column=2).value
title = str(title_cell.value)

if not os.path.exists(output_file):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{title}_{Model}"  # 设置第一个表单的标题
    ws["A1"] = "Model"
    ws["B1"] = "Question"
    ws["C1"] = "Model Answer"
    ws['D1'] = "rect Answer"
    wb.save(output_file)
else:
    wb = openpyxl.load_workbook(output_file)
    ws = wb.active

def create_query(current_row):
    question_cell = input_ws.cell(row=current_row, column=3)
    if question_cell.value == None:
        return None
    question = re.sub(r'\n', ' ', str(question_cell.value))
    return question

def chat_response(query):
    model = ChatOpenAI(temperature=0)

    prompt = ChatPromptTemplate.from_messages([
        SystemMessagePromptTemplate.from_template("You must answer the question based on A or B or C or D"),
        # MessagesPlaceholder(variable_name="chat_history"),
        HumanMessagePromptTemplate.from_template("{input}"),
    ])

    memory = ConversationSummaryBufferMemory(llm=model,memory_key="chat_history",return_messages=True,k=3,max_token_limit=50)

    # print(memory.prompt.template)

    # chain = ConversationChain(
    #     llm=model,
    #     # verbose=True,
    #     prompt=prompt,
    #     memory=memory
    # )

    chain = prompt | model

    response = chain.invoke(input={
        "input": query,
    })
    return response

def get_usage_tokens(query):
    response_metadata = chat_response(query).response_metadata['token_usage']['total_tokens']
    return response_metadata

total_tokens = 0
i = 2   
if __name__ == "__main__":
    while True:
        query = create_query(current_row)
        if query == None:
            sheet_index += 1
            current_row = 2
            i = 2
            ws = wb.create_sheet(title=f"{title}")
            ws["A1"] = "Model"
            ws["B1"] = "Question"
            ws["C1"] = "Model Answer"
            ws['D1'] = "rect Answer"
            if sheet_index >= len(input_wb.worksheets):
                break
            input_ws = input_wb.worksheets[sheet_index]
            continue

        response = chat_response(query).content
        content = re.search('[A-D]',response)

        Model_cell = input_ws.cell(row=current_row, column=2)
        question_cell = input_ws.cell(row=current_row, column=3)
        rect_answer_cell = input_ws.cell(row=current_row, column=6)

        if content:
            model = ws.cell(row=current_row, column=1, value=Model_cell.value)
            question = ws.cell(row=current_row, column=2, value=question_cell.value)
            result = ws.cell(row=ws.max_row,column=3,value=content.group())
            rect_answer = ws.cell(row=current_row, column=4, value=rect_answer_cell.value)

            print(f"Chat_BOT: {response}")
            wb.save(output_file)
            token = get_usage_tokens(query)
            print(f"第{i}题已成功存储至 Excel 文件中，结果为：{response},token消耗为：{token}")
            current_row += 1
            i+=1
            total_tokens += token
    print(f"total_tokens: {total_tokens}")