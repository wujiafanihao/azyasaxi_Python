import requests
from bs4 import BeautifulSoup
import openpyxl
import os
import re

input_file = "douban_video_id.xlsx"
output_file = "douban_actor.xlsx"
input_wb = openpyxl.load_workbook(input_file)
input_ws = input_wb.active

if not os.path.exists(output_file):
    wb = openpyxl.Workbook()
    ws = wb.active
    wb.save(output_file)
else:
    wb = openpyxl.load_workbook(output_file)
    ws = wb.active

current_row = 1

def get_next_id():
    global current_row
    id = input_ws.cell(row=current_row, column=1).value
    if id is None:
        return None
    current_row += 1
    return id

def extract_data(id):
    url = f'https://movie.douban.com/subject/{id}/celebrities'
    headers = {
        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/58.0.3029.110 Safari/537.36'
    }
    response = requests.get(url, headers=headers)

    if response.status_code == 200:
        soup = BeautifulSoup(response.text, 'html.parser')
        target_elements = soup.select('#celebrities > div:nth-of-type(2) ul li div span:nth-of-type(1) a')
        span_elements = soup.select('#celebrities > div:nth-of-type(2) ul li div span:nth-of-type(2)')
        title_element = soup.select_one('#content > h1')  

        movie_title = title_element.get_text(strip=True) if title_element else None
        # 使用正则表达式提取电影名称
        movie_name = re.search(r'^(.+?)\s*的全部演职员', movie_title)
        if movie_name:
            movie_name = movie_name.group(1)

        actor_data = []
        for element, span in zip(target_elements, span_elements):
            if len(actor_data) >= 2:  # 如果已经获取了两个演员信息，则跳出循环
                break

            actor_name = element.get_text()
            # 使用正则表达式过滤掉非中文字符
            chinese_actor_name = re.sub(r'[^\u4e00-\u9fa5]', '', actor_name)
            text = span.get_text()
            if re.search(r'^自己$', text):
                actor_data.append((chinese_actor_name, ""))
            # 使用正则表达式提取 (饰 范闲) 部分，包括括号和“饰”字
            role = re.search(r'\(饰\s*(.+)\)', text)
            if role:
                role_text = role.group(0)  
                actor_data.append((chinese_actor_name, role_text))

        actor_info = '-->'.join([f"{actor_name}{role_text}" for actor_name, role_text in actor_data])
        print(f"《{movie_name}》-->{actor_info}")

        store_to_excel(movie_name, actor_info)

        return actor_data
    else:
        print(f"Failed to retrieve the page for ID {id}, status code: {response.status_code}")
        return []

def store_to_excel(movie_name, actor_info):
    ws.append([movie_name, actor_info])
    wb.save(output_file)

id = get_next_id()
while id is not None:
    data = extract_data(id)
    id = get_next_id()