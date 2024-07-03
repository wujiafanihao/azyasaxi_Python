import urllib.parse
import os
import pandas as pd
from openpyxl import load_workbook
from openpyxl import Workbook
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.chrome.options import Options
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
import chardet
import re

txt_name = 'star.txt'
output_dir = 'Actor.xlsx'

def translate_ASCII(name):
    print(urllib.parse.quote(name))
    return urllib.parse.quote(name)

def detect_encoding(file_name):
    with open(file_name, 'rb') as file:
        raw_data = file.read()
    print(chardet.detect(raw_data))
    return chardet.detect(raw_data)['encoding']

encoding = detect_encoding(txt_name)

def get_names(file_name):
    try:
        with open(file_name, 'r', encoding=encoding) as file:
            names = file.readlines()
        return [name.strip() for name in names if name.strip()]
    except FileNotFoundError:
        print(f"文件 {file_name} 未找到。")
        return None
    except Exception as e:
        print(f"发生错误: {e}")
        return None

file_name = txt_name
names = get_names(file_name)

chrome_options = Options()
chrome_options.add_experimental_option("prefs", {
    "profile.managed_default_content_settings.images": 2  
})
chrome_options.add_argument('--enable-chrome-browser-cloud-management')
chrome_options.binary_location = "E:/python/chrome-win64/chrome.exe"  
driver_path = "E:/python/chromedriver.exe"  

service = Service(driver_path)
driver = webdriver.Chrome(service=service, options=chrome_options)
wait = WebDriverWait(driver, 10)

def get_tvs_and_movies():
    movie_list = set()
    
    try:
        # Extract movies and TV shows
        movie_elements = wait.until(EC.presence_of_all_elements_located(
            (By.CSS_SELECTOR, 'div.movieAndTvPosterWrapper__XfGC span.text_RUbYG[data-text="true"]:not(dl.descWrapper_rP_ci span.text_RUbYG)')
        ))

        for element in movie_elements:
            # Get text directly inside span
            span_text = element.text.strip()
            if span_text and not re.match(r'\d{4}(-\d{1,2}(-\d{1,2})?)?', span_text):
                movie_list.add(span_text)

            # Get text inside the nested a tag if present
            try:
                a_tag = element.find_element(By.CSS_SELECTOR, 'a.innerLink_yDkYh')
                a_text = a_tag.text.strip()
                if a_text and not re.match(r'\d{4}(-\d{1,2}(-\d{1,2})?)?', a_text):
                    movie_list.add(a_text)
            except:
                # Ignore if no <a> tag is found
                pass

    except Exception as e:
        print(f"Error extracting movie and TV show names: {e}")

    print(list(movie_list))
    return list(movie_list)

def create_excel_file(file_name):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    workbook.save(file_name)

def write_to_excel(name, movie_list, start_row):
    data = {
        "A": [name],
    }

    for i, tv in enumerate(movie_list):
        data[chr(67 + i)] = [tv]  

    df = pd.DataFrame(data)

    file_name = output_dir
    if not os.path.exists(file_name):
        create_excel_file(file_name)
        start_row = 0  

    try:
        with pd.ExcelWriter(file_name, engine='openpyxl', mode='a', if_sheet_exists='overlay') as writer:
            df.to_excel(writer, index=False, header=False, startrow=start_row)

    except Exception as e:
        print(f"写入Excel文件时出错: {e}")

def get_start_row(file_name):
    if os.path.exists(file_name):
        workbook = load_workbook(file_name)
        sheet = workbook.active
        return sheet.max_row
    return 0

start_row = get_start_row(output_dir)

for name in names:
    try:
        encoded_name = translate_ASCII(name)
        url = f"https://baike.baidu.com/item/{encoded_name}"
        driver.get(url)
        movie_list = get_tvs_and_movies()
        if movie_list:
            write_to_excel(name, movie_list, start_row)
            print(f"movie_list：{movie_list}")
            movie_list.clear()
            start_row += 1  # Increment the row for the next entry
    except Exception as e:
        print(f"处理 {name} 时出错: {e}")

driver.quit()
